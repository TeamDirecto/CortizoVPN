import json
import os

from openpyxl import load_workbook

from app.db import run_readonly_query


DEFAULT_FILE = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "data",
    "Cortizo_rangos_100_desde_161001.xlsx",
)
SHEET_NAME = "Asignación final"
GROUP_DEFAULTS_FILE = "/etc/vici-users/group_defaults.json"
GROUP_TEMPLATES_FILE = "/etc/vici-users/group_templates.json"

USER_CLONE_FIELDS = [
    "delete_users", "delete_user_groups", "delete_lists", "delete_campaigns",
    "delete_ingroups", "delete_remote_agents", "load_leads", "campaign_detail",
    "ast_admin_access", "ast_delete_phones", "delete_scripts", "modify_leads",
    "hotkeys_active", "change_agent_campaign", "agent_choose_ingroups",
    "closer_campaigns", "scheduled_callbacks", "agentonly_callbacks",
    "agentcall_manual", "vicidial_recording", "vicidial_transfers",
    "delete_filters", "alter_agent_interface_options", "closer_default_blended",
    "delete_call_times", "modify_call_times", "modify_users", "modify_campaigns",
    "modify_lists", "modify_scripts", "modify_filters", "modify_ingroups",
    "modify_usergroups", "modify_remoteagents", "modify_servers", "view_reports",
    "vicidial_recording_override", "alter_custdata_override", "qc_enabled",
    "qc_user_level", "qc_pass", "qc_finish", "qc_commit", "add_timeclock_log",
    "modify_timeclock_log", "delete_timeclock_log", "alter_custphone_override",
    "vdc_agent_api_access", "modify_inbound_dids", "delete_inbound_dids",
    "alert_enabled", "download_lists", "agent_shift_enforcement_override",
    "manager_shift_enforcement_override", "shift_override_flag", "export_reports",
    "delete_from_dnc", "allow_alerts", "agent_choose_territories",
    "agent_call_log_view_override", "callcard_admin", "agent_choose_blended",
    "realtime_block_user_info", "custom_fields_modify", "force_change_password",
    "agent_lead_search_override", "modify_shifts", "modify_phones",
    "modify_carriers", "modify_labels", "modify_statuses", "modify_voicemail",
    "modify_audiostore", "modify_moh", "modify_tts", "preset_contact_search",
    "modify_contacts", "modify_same_user_level", "admin_hide_lead_data",
    "admin_hide_phone_data", "agentcall_email", "modify_email_accounts",
    "alter_admin_interface_options", "max_inbound_calls",
    "modify_custom_dialplans", "wrapup_seconds_override", "modify_languages",
    "selected_language", "user_choose_language", "ignore_group_on_search",
    "api_list_restrict", "api_allowed_functions", "lead_filter_id",
    "admin_cf_show_hidden", "agentcall_chat", "user_hide_realtime",
    "access_recordings", "modify_colors", "user_new_lead_limit", "api_only_user",
    "modify_auto_reports", "modify_ip_lists", "ignore_ip_list",
    "ready_max_logout", "export_gdpr_leads", "pause_code_approval",
    "max_hopper_calls", "max_hopper_calls_hour", "mute_recordings",
    "hide_call_log_info", "next_dial_my_callbacks", "user_admin_redirect_url",
    "max_inbound_filter_enabled", "max_inbound_filter_statuses",
    "max_inbound_filter_ingroups", "max_inbound_filter_min_sec",
    "status_group_id", "two_factor_override", "manual_dial_filter",
    "download_invalid_files", "user_group_two", "modify_dial_prefix",
    "inbound_credits", "hci_enabled", "manual_dial_lead_id",
]


def _load_json(path):
    with open(path) as handle:
        data = json.load(handle)
    if not isinstance(data, dict):
        raise RuntimeError("JSON invalido: {0}".format(path))
    return data


def load_group_defaults():
    return _load_json(GROUP_DEFAULTS_FILE)


def load_group_templates():
    return _load_json(GROUP_TEMPLATES_FILE)


def _read_targets(path=DEFAULT_FILE):
    if not os.path.exists(path):
        raise RuntimeError("No existe el archivo de migracion: {0}".format(path))

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise RuntimeError("No existe la hoja '{0}'".format(SHEET_NAME))
        ws = wb[SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
        header = [str(v).strip() if v is not None else "" for v in rows[0]]
        idx = {name: i for i, name in enumerate(header)}

        required = ["User Group", "User", "Nombre", "Nueva extensión"]
        missing = [name for name in required if name not in idx]
        if missing:
            raise RuntimeError("Faltan columnas: {0}".format(", ".join(missing)))

        targets = []
        for row in rows[1:]:
            user = row[idx["User"]]
            if user in (None, ""):
                continue
            extension = row[idx["Nueva extensión"]]
            ext = str(int(extension)) if isinstance(extension, float) else str(extension).strip()
            targets.append({
                "user": str(user).strip(),
                "full_name": str(row[idx["Nombre"]] or "").strip(),
                "user_group": str(row[idx["User Group"]] or "").strip(),
                "extension": ext,
            })
        return targets
    finally:
        wb.close()


def _schema(config):
    rows = run_readonly_query(config, "SHOW COLUMNS FROM vicidial_users", "master")
    return [row[0] for row in rows if row]


def _existing_users(config):
    rows = run_readonly_query(
        config,
        "SELECT user, full_name, user_group, phone_login, phone_pass, active, user_level "
        "FROM vicidial_users ORDER BY user",
        "master",
    )
    result = {}
    for row in rows:
        result[row[0]] = {
            "user": row[0],
            "full_name": row[1],
            "user_group": row[2],
            "phone_login": row[3],
            "phone_pass": row[4],
            "active": row[5],
            "user_level": row[6],
        }
    return result


def _template_status(config, groups, templates):
    if not groups:
        return {}
    quoted = ", ".join("'" + g.replace("'", "''") + "'" for g in groups)
    rows = run_readonly_query(
        config,
        "SELECT user, user_group, active, user_level FROM vicidial_users "
        "WHERE user_group IN ({0}) ORDER BY user_group, user".format(quoted),
        "master",
    )
    by_key = {(row[0], row[1]): row for row in rows}
    status = {}
    for group in groups:
        template = templates.get(group)
        row = by_key.get((template, group)) if template else None
        status[group] = {
            "template_user": template,
            "exists": bool(row),
            "active": row[2] if row else None,
            "user_level": row[3] if row else None,
        }
    return status


def build_agent_plan(config, path=DEFAULT_FILE):
    targets = _read_targets(path)
    schema = _schema(config)
    existing = _existing_users(config)
    defaults = load_group_defaults()
    templates = load_group_templates()

    groups = sorted(set(item["user_group"] for item in targets))
    template_status = _template_status(config, groups, templates)

    missing_schema = [
        field for field in USER_CLONE_FIELDS
        if field not in schema
    ]
    required_identity = [
        "user", "pass", "full_name", "user_level", "user_group",
        "active", "phone_login", "phone_pass",
    ]
    missing_identity_schema = [
        field for field in required_identity if field not in schema
    ]
    missing_defaults = [group for group in groups if group not in defaults]
    missing_templates = [
        group for group in groups
        if not template_status.get(group, {}).get("exists")
    ]

    duplicate_users = []
    duplicate_extensions = []
    seen_users = set()
    seen_extensions = set()
    for target in targets:
        if target["user"] in seen_users:
            duplicate_users.append(target["user"])
        seen_users.add(target["user"])
        if target["extension"] in seen_extensions:
            duplicate_extensions.append(target["extension"])
        seen_extensions.add(target["extension"])

    items = []
    create_count = 0
    update_count = 0
    correct_count = 0

    for target in targets:
        current = existing.get(target["user"])
        desired = {
            "user": target["user"],
            "full_name": target["full_name"],
            "user_group": target["user_group"],
            "phone_login": target["extension"],
            "phone_pass": target["extension"],
            "active": "Y",
            "user_level": "1",
            "template_user": templates.get(target["user_group"]),
            "password_configured": target["user_group"] in defaults,
        }

        if current is None:
            status = "CREATE"
            create_count += 1
        else:
            mismatched = [
                key for key in (
                    "full_name", "user_group", "phone_login",
                    "phone_pass", "active", "user_level",
                )
                if str(current.get(key) or "") != str(desired.get(key) or "")
            ]
            if mismatched:
                status = "UPDATE"
                update_count += 1
            else:
                status = "EXISTS_CORRECT"
                correct_count += 1

        items.append({
            "status": status,
            "user": target["user"],
            "full_name": target["full_name"],
            "user_group": target["user_group"],
            "extension": target["extension"],
            "template_user": templates.get(target["user_group"]),
            "password_configured": target["user_group"] in defaults,
            "current": current,
        })

    blockers = []
    if duplicate_users:
        blockers.append("DUPLICATE_USERS")
    if duplicate_extensions:
        blockers.append("DUPLICATE_EXTENSIONS")
    if missing_identity_schema:
        blockers.append("IDENTITY_SCHEMA_MISMATCH")
    # Diferencias de version entre EHECTO y CortizoVPN son esperables.
    # Solo se clonan los campos de la allowlist que existen en el esquema destino.
    if missing_defaults:
        blockers.append("DEFAULT_PASSWORD_NOT_CONFIGURED")
    if missing_templates:
        blockers.append("USER_TEMPLATE_ROW_MISSING")

    return {
        "source_file": path,
        "sheet": SHEET_NAME,
        "target_count": len(targets),
        "group_count": len(groups),
        "create_count": create_count,
        "update_count": update_count,
        "correct_count": correct_count,
        "duplicate_users": sorted(set(duplicate_users)),
        "duplicate_extensions": sorted(set(duplicate_extensions)),
        "missing_identity_schema": missing_identity_schema,
        "missing_clone_schema": missing_schema,
        "supported_clone_fields": [
            field for field in USER_CLONE_FIELDS if field in schema
        ],
        "missing_defaults": missing_defaults,
        "missing_templates": missing_templates,
        "template_status": template_status,
        "clone_fields_count": len(USER_CLONE_FIELDS),
        "blockers": blockers,
        "write_ready": not blockers,
        "items": items,
    }
